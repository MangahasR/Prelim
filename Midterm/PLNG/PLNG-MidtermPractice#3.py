#31
import pymysql.cursors

connection = pymysql.connect(
    host='localhost',
    user='root',
    password='',
    db='students',
    charset='utf8mb4',
    cursorclass=pymysql.cursors.DictCursor
)

try:
    with connection.cursor() as cursor:
        sql = "SELECT `id`, `firstname`, `lastname` FROM `users`"
        cursor.execute(sql)

        for row in cursor.fetchall():
            firstname = str(row["firstname"])
            lastname = str(row["lastname"])

            print("firstname : " + firstname)
            print("lastname : " + lastname)

finally:
    connection.close()

#32
import tkinter as tk

window = tk.Tk()
window.title("My First GUI")
window.geometry("300x200")

window.mainloop()

#33
import tkinter as tk

window = tk.Tk()
window.title("Entry Form")

label = tk.Label(window, text="Enter your name:")
label.pack()

entry = tk.Entry(window)
entry.pack()

window.mainloop()

#34
import tkinter as tk

window = tk.Tk()
window.title("Listbox")

listbox = tk.Listbox(window)
listbox.pack()

listbox.insert(tk.END, "Apple")
listbox.insert(tk.END, "Banana")
listbox.insert(tk.END, "Orange")

window.mainloop()


#35
import tkinter as tk
from openpyxl import Workbook, load_workbook
import os

filename = "registration.xlsx"

if os.path.exists(filename):
    workbook = load_workbook(filename)
    sheet = workbook.active
else:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Age", "Email"])

def register():
    name = name_entry.get()
    age = age_entry.get()
    email = email_entry.get()

    sheet.append([name, age, email])
    workbook.save(filename)

    name_entry.delete(0, tk.END)
    age_entry.delete(0, tk.END)
    email_entry.delete(0, tk.END)

window = tk.Tk()
window.title("Registration Form")

tk.Label(window, text="Name").pack()
name_entry = tk.Entry(window)
name_entry.pack()

tk.Label(window, text="Age").pack()
age_entry = tk.Entry(window)
age_entry.pack()

tk.Label(window, text="Email").pack()
email_entry = tk.Entry(window)
email_entry.pack()

tk.Button(window, text="Register", command=register).pack()

window.mainloop()


#36
from random import randint

rand = randint(1, 100)
counter = 0

while True:
    counter += 1
    num = int(input("Enter values between 1 and 100 (0 exit):"))

    if num == 0:
        print("You canceled the game.")
        break
    elif num < rand:
        print("Enter a Higher Number.")
        continue
    elif num > rand:
        print("Enter a Lower Number.")
        continue
    else:
        print("Randomly selected number {0}!".format(rand))
        print("Your guess number {0}".format(counter))


#37
favorite_fruits = ['apples', 'erics', 'oranges']

print("My favourite fruits {}".format(favorite_fruits))


#38
list = [3.14, 2.54, "CPU", "WINDOWS 10"]

print(list)


#39
weekdays = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

print(weekdays[4])


#40
list = [3, 1, 2]

list[0] = 1
list[1] = 1
list[2] = 2

print(list)


#41
weekdays = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

print("Friday" in weekdays)
print("Saturday" in weekdays)


#42
list1 = [10, 21, 4, 45, 66, 93]

for num in list1:

    # checking condition
    if num % 2 == 0:
        print(num, end=" ")


#43
numbers = [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18]
numbers2 = [21, 22, 23, 24, 25]
combine = numbers + numbers2

for number in combine:
    if sayi % 4 == 0:
        print(number)


#44
i = 1

while i <= 12:
    print(" {}. class".format(i))
    i += 1


#45
number = input("enter a number: ")
sum = 0

for numberofdigit in number:
    sum += int(numberofdigit)

print("the sum of the digits of the number:", sum)


#46
number = input("enter an number: ")
carpim = 1

for digitofnumber in str(number):
    if int(digitofnumber) != 0:
        multiplication *= int(digitofnumber)

print("digits of the number: ", multiplication)


#47
text = input("enter a name: ")

for i in range(10):
    print(text)


#48
text = input("enter a name: ")

i = 0

while i < 10:
    i += 1
    print(text)


#49
list = [35, 48, 52, 57, 68, 84]

total = sum(list)
piece = len(list)

print(total / piece)


#50
total = 0

while True:
    num = float(input("enter a number: "))

    if num == 0:
        break

    total += num

print("The sum of the numbers: ", total)



